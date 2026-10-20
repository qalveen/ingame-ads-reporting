"""
Publishes the monthly A&S Inventory Report (Excel + chart PNGs) from the mart tables.
Called by run_pipeline.py after all error-severity data tests pass.
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
import matplotlib.ticker as mtick
import pandas as pd
import yaml
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]

# ---- style -------------------------------------------------------------------------
FONT = "Arial"
NAVY = "1F2A44"
GRID = "D9DDE3"
PALETTE = ["#2F6DB5", "#E07B39", "#3A9E6F", "#8E5EA2", "#C94B4B", "#7A8899"]
H_FILL = PatternFill("solid", fgColor=NAVY)
H_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
B_FONT = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
TITLE = Font(name=FONT, size=16, bold=True, color=NAVY)
SUB = Font(name=FONT, size=10, italic=True, color="5F6B7A")
THIN = Border(bottom=Side(style="thin", color=GRID))
RED_FILL = PatternFill("solid", fgColor="F8D7D7")
AMBER_FILL = PatternFill("solid", fgColor="FCEBC8")
GREEN_FILL = PatternFill("solid", fgColor="D8EFDF")
GREY_FILL = PatternFill("solid", fgColor="ECEFF3")

FMT = {
    "int": "#,##0;(#,##0);-",
    "pct": "0.0%;(0.0%);-",
    "pts": "+0.0%;-0.0%;-",
    "usd": "$#,##0.00;($#,##0.00);-",
    "dec": "0.00",
    "date": "yyyy-mm-dd",
}


def write_table(ws, df: pd.DataFrame, start_row: int, start_col: int = 1,
                formats: dict | None = None, headers: dict | None = None) -> tuple[int, int]:
    """Write a dataframe as a formatted table. Returns (first_data_row, last_data_row)."""
    formats, headers = formats or {}, headers or {}
    for j, col in enumerate(df.columns):
        c = ws.cell(row=start_row, column=start_col + j, value=headers.get(col, col.replace("_", " ").title()))
        c.fill, c.font = H_FILL, H_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[start_row].height = 30
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, (col, val) in enumerate(zip(df.columns, row)):
            if isinstance(val, float) and pd.isna(val):
                val = None
            elif hasattr(val, "item"):
                val = val.item()
            c = ws.cell(row=start_row + i, column=start_col + j, value=val)
            c.font, c.border = B_FONT, THIN
            if col in formats:
                c.number_format = FMT[formats[col]]
    return start_row + 1, start_row + len(df)


def autosize(ws, min_w: int = 9, max_w: int = 42) -> None:
    widths = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is None or c.font is None or c.font.size == 16:
                continue
            L = len(str(c.value)) if not isinstance(c.value, float) else 10
            if c.row and ws.row_dimensions[c.row].height == 30:   # wrapped headers
                L = min(L, 14)
            widths[c.column_letter] = max(widths.get(c.column_letter, 0), L)
    for col, w in widths.items():
        ws.column_dimensions[col].width = max(min_w, min(max_w, w + 2))


def header_block(ws, title: str, subtitle: str) -> None:
    ws["A1"] = title
    ws["A1"].font = TITLE
    ws["A2"] = subtitle
    ws["A2"].font = SUB
    ws.sheet_view.showGridLines = False


# ---- charts ------------------------------------------------------------------------
def _style(ax):
    for s in ["top", "right"]:
        ax.spines[s].set_visible(False)
    ax.spines["left"].set_color("#9AA5B1")
    ax.spines["bottom"].set_color("#9AA5B1")
    ax.tick_params(colors="#3D4752", labelsize=9)
    ax.grid(axis="y", color="#E6E9ED", linewidth=0.8)
    ax.set_axisbelow(True)


def make_charts(daily: pd.DataFrame, placements: pd.DataFrame, campaigns: pd.DataFrame,
                titles: pd.DataFrame, out: Path) -> dict:
    out.mkdir(parents=True, exist_ok=True)
    plt.rcParams["font.family"] = "DejaVu Sans"
    paths = {}
    names = dict(zip(titles["title_id"], titles["title_name"]))

    # 1. daily paid impressions (7-day avg) by title
    fig, ax = plt.subplots(figsize=(9, 3.6), dpi=140)
    for k, (tid, g) in enumerate(daily.groupby("title_id")):
        ax.plot(g["event_date"], g["paid_impressions"], color=PALETTE[k], alpha=0.25, linewidth=1)
        ax.plot(g["event_date"], g["paid_impressions_7d_avg"], color=PALETTE[k], linewidth=2.2, label=names[tid])
    _style(ax)
    ax.yaxis.set_major_formatter(mtick.FuncFormatter(lambda v, _: f"{v/1000:.0f}K"))
    ax.set_title("Paid impressions per day by title (line = 7-day average)", loc="left", fontsize=11, color="#1F2A44")
    ax.legend(frameon=False, fontsize=8, ncol=4, loc="upper left")
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%b %d"))
    fig.tight_layout()
    paths["daily"] = out / "daily_paid_impressions.png"
    fig.savefig(paths["daily"]); plt.close(fig)

    # 2. viewability by placement vs 70% benchmark
    p = placements.dropna(subset=["viewability_rate"]).sort_values("viewability_rate")
    fig, ax = plt.subplots(figsize=(9, 4), dpi=140)
    colors = [PALETTE[0] if v >= 0.7 else PALETTE[4] for v in p["viewability_rate"]]
    ax.barh(p["placement_name"] + " (" + p["title_id"] + ")", p["viewability_rate"], color=colors)
    ax.axvline(0.7, color="#3D4752", linestyle="--", linewidth=1)
    ax.text(0.705, -0.6, "70% target", fontsize=8, color="#3D4752")
    _style(ax); ax.grid(axis="x", color="#E6E9ED"); ax.grid(axis="y", visible=False)
    ax.xaxis.set_major_formatter(mtick.PercentFormatter(1.0))
    ax.set_xlim(0, 1)
    ax.set_title("Viewability rate by placement (MRC standard)", loc="left", fontsize=11, color="#1F2A44")
    fig.tight_layout()
    paths["viewability"] = out / "viewability_by_placement.png"
    fig.savefig(paths["viewability"]); plt.close(fig)

    # 3. flight-to-date pacing by campaign with tolerance band
    c = campaigns.sort_values("pacing_pct_to_date")
    fig, ax = plt.subplots(figsize=(9, 4.6), dpi=140)
    col = [PALETTE[4] if s in ("UNDER-PACING", "NO DELIVERY") else PALETTE[1] if s == "OVER-PACING" else PALETTE[2]
           for s in c["pacing_status"]]
    vals = c["pacing_pct_to_date"].fillna(0)
    ax.barh(c["campaign_id"] + "  " + c["advertiser"], vals, color=col)
    for y, (v, s) in enumerate(zip(vals, c["pacing_status"])):
        ax.text(v + 0.01, y, "No delivery" if s == "NO DELIVERY" else f"{v:.0%}", va="center", fontsize=7.5, color="#3D4752")
    ax.axvspan(0.9, 1.1, color="#3A9E6F", alpha=0.08)
    ax.set_xlim(0, max(1.3, vals.max() + 0.12))
    ax.axvline(1.0, color="#3D4752", linewidth=1)
    _style(ax); ax.grid(axis="x", color="#E6E9ED"); ax.grid(axis="y", visible=False)
    ax.xaxis.set_major_formatter(mtick.PercentFormatter(1.0))
    ax.set_title("Flight-to-date pacing vs goal (shaded = 90–110% tolerance)", loc="left", fontsize=11, color="#1F2A44")
    fig.tight_layout()
    paths["pacing"] = out / "campaign_pacing.png"
    fig.savefig(paths["pacing"]); plt.close(fig)
    return paths


# ---- narrative ---------------------------------------------------------------------
def callouts(titles: pd.DataFrame, campaigns: pd.DataFrame, placements: pd.DataFrame,
             dq: pd.DataFrame, tests: list[dict]) -> list[str]:
    out = []
    tot_paid = titles["paid_impressions"].sum()
    tot_rev = titles["revenue_usd"].sum()
    top = titles.sort_values("revenue_usd", ascending=False).iloc[0]
    mom = f"{top['revenue_mom']:+.0%} MoM" if pd.notna(top["revenue_mom"]) else "no prior month on file"
    out.append(f"{tot_paid:,.0f} paid impressions and ${tot_rev:,.0f} billable revenue across "
               f"{len(titles)} titles; {top['title_name']} led revenue (${top['revenue_usd']:,.0f}, {mom}).")
    low_v = placements[placements["viewability_rate"] < 0.7]
    if len(low_v):
        names = ", ".join(f"{r.placement_name} ({r.title_id}, {r.viewability_rate:.2%})" for r in low_v.itertuples())
        out.append(f"Below 70% viewability: {names}.")
    risk = campaigns[campaigns["pacing_status"].isin(["UNDER-PACING", "NO DELIVERY"])]
    for r in risk.itertuples():
        if r.pacing_status == "NO DELIVERY":
            out.append(f"{r.campaign_id} {r.advertiser}: booked {r.booked_impressions:,.0f} but no delivery — "
                       f"confirm trafficking with AdOps ({r.adops_note or 'no note'}).")
        else:
            gap = r.expected_impressions_to_date - r.delivered_to_date
            out.append(f"{r.campaign_id} {r.advertiser} under-pacing at {r.pacing_pct_to_date:.0%} "
                       f"flight-to-date ({gap:,.0f} impressions behind).")
    over = campaigns[campaigns["pacing_status"] == "OVER-PACING"]
    if len(over):
        out.append("Over-pacing (>110%): " + ", ".join(f"{r.campaign_id} ({r.pacing_pct_to_date:.0%})"
                                                      for r in over.itertuples())
                   + ". Over-delivery is not billed; review frequency caps / goal settings with AdOps.")
    new = placements[placements["is_new_inventory"]]
    for r in new.itertuples():
        out.append(f"New inventory: {r.placement_name} ({r.title_id}) launched {r.launch_date:%b %d}; "
                   f"{r.paid_impressions:,.0f} paid impressions, {r.sell_through_rate:.0%} sell-through, "
                   f"{r.viewability_rate:.0%} viewability.")
    flags = dq[dq["status"] == "flag"]
    if len(flags):
        out.append("Data-quality flags: " + "; ".join(f"{r.check_name} = {r.value:,.0f}" for r in flags.itertuples())
                   + ". Excluded from reported totals; see Data Quality tab.")
    warns = [t for t in tests if t["status"] == "WARN"]
    out.append(f"{sum(t['status'] == 'PASS' for t in tests)} of {len(tests)} automated checks passed; "
               f"{len(warns)} warnings need follow-up, 0 blocking failures.")
    return out


# ---- workbook ----------------------------------------------------------------------
def build_report(con, month: str, tests: list[dict], out_dir: Path) -> Path:
    q = lambda sql: con.execute(sql).df()
    titles = q("select * from mart_title_monthly")
    campaigns = q("select * from mart_campaign_monthly")
    placements = q("select * from mart_placement_monthly")
    freq = q("select * from mart_frequency_distribution")
    daily = q("select * from mart_daily_title")
    dq = q("select * from mart_dq_summary")
    scope = q("select * from stg_adops_scope order by campaign_id")
    kpis = pd.DataFrame(yaml.safe_load((ROOT / "config/kpi_definitions.yml").read_text())["kpis"])
    month_label = pd.Timestamp(f"{month}-01").strftime("%B %Y")
    run_at = datetime.now().strftime("%Y-%m-%d %H:%M")

    charts = make_charts(daily, placements, campaigns, titles, out_dir / "charts")
    wb = Workbook()

    # Summary ----------------------------------------------------------------------
    ws = wb.active
    ws.title = "Summary"
    header_block(ws, f"Ads & Sponsorships Inventory Report — {month_label}",
                 f"Custom in-game inventory, AdOps-approved scope ({int(scope['include_in_report'].sum())} campaigns). "
                 f"Generated {run_at}. Synthetic portfolio data.")
    tot = dict(
        ad_requests=titles["ad_requests"].sum(),
        paid=titles["paid_impressions"].sum(),
        rev=titles["revenue_usd"].sum(),
    )
    viewable = (placements["viewability_rate"].fillna(0) * placements["paid_impressions"]).sum()
    tiles = [
        ("Ad requests", tot["ad_requests"], "int"),
        ("Paid impressions", tot["paid"], "int"),
        ("Sell-through", tot["paid"] / tot["ad_requests"], "pct"),
        ("Viewability", viewable / tot["paid"], "pct"),
        ("Revenue (USD)", tot["rev"], "usd"),
        ("eCPM (USD)", tot["rev"] / tot["paid"] * 1000, "usd"),
    ]
    for j, (label, val, f) in enumerate(tiles):
        col = 1 + j * 2
        ws.cell(row=4, column=col, value=label).font = Font(name=FONT, size=9, color="5F6B7A")
        c = ws.cell(row=5, column=col, value=float(val))
        c.font = Font(name=FONT, size=15, bold=True, color=NAVY)
        c.number_format = FMT[f].split(";")[0]

    ws["A7"] = "Key callouts"; ws["A7"].font = BOLD
    notes = callouts(titles, campaigns, placements, dq, tests)
    for i, n in enumerate(notes):
        c = ws.cell(row=8 + i, column=1, value=f"•  {n}")
        c.font = B_FONT
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=8 + i, start_column=1, end_row=8 + i, end_column=12)
        ws.row_dimensions[8 + i].height = 28
    r = 9 + len(notes)
    ws.cell(row=r, column=1, value="By title").font = BOLD
    tcols = ["title_name", "ad_requests", "fill_rate", "paid_impressions", "sell_through_rate", "viewability_rate",
             "paid_unique_reach", "avg_frequency", "revenue_usd", "ecpm_usd", "paid_impressions_mom", "revenue_mom"]
    write_table(ws, titles[tcols], r + 1,
                formats=dict(ad_requests="int", fill_rate="pct", paid_impressions="int", sell_through_rate="pct",
                             viewability_rate="pct", paid_unique_reach="int", avg_frequency="dec",
                             revenue_usd="usd", ecpm_usd="usd", paid_impressions_mom="pts", revenue_mom="pts"),
                headers=dict(title_name="Title", paid_unique_reach="Unique Reach", ecpm_usd="eCPM (USD)",
                             revenue_usd="Revenue (USD)", paid_impressions_mom="Paid Imps MoM", revenue_mom="Revenue MoM"))
    autosize(ws)
    ws.column_dimensions["A"].width = 20
    img = XLImage(str(charts["daily"])); img.width, img.height = 900, 360
    ws.add_image(img, f"A{r + 8}")

    # Title KPIs -------------------------------------------------------------------
    ws = wb.create_sheet("Title KPIs")
    header_block(ws, "Title KPIs", f"{month_label} vs prior month. MoM rates shown as % change; rate KPIs as point change.")
    write_table(ws, titles, 4, formats=dict(
        ad_requests="int", fill_rate="pct", paid_impressions="int", house_impressions="int", sell_through_rate="pct",
        viewability_rate="pct", paid_unique_reach="int", avg_frequency="dec", live_campaigns="int",
        revenue_usd="usd", ecpm_usd="usd", paid_impressions_mom="pts", revenue_mom="pts",
        viewability_mom_pts="pts", sell_through_mom_pts="pts"))
    ws.freeze_panes = "C5"
    autosize(ws)

    # Campaign Detail --------------------------------------------------------------
    ws = wb.create_sheet("Campaign Detail")
    header_block(ws, "Campaign Detail", "In-scope campaigns only. Pacing status uses flight-to-date delivery vs straight-line goal.")
    ccols = ["campaign_id", "advertiser", "vertical", "title_id", "placement_ids", "sold_by", "flight_start", "flight_end",
             "booked_impressions", "expected_impressions_month", "impressions", "delivery_pct_month",
             "delivered_to_date", "expected_impressions_to_date", "pacing_pct_to_date", "pacing_status",
             "viewable_impressions", "viewability_rate", "unique_reach", "avg_frequency", "video_starts",
             "video_completion_rate", "billable_impressions", "cpm_usd", "revenue_usd", "adops_note"]
    first, last = write_table(ws, campaigns[ccols], 4, formats=dict(
        flight_start="date", flight_end="date", booked_impressions="int", expected_impressions_month="int",
        impressions="int", delivery_pct_month="pct", delivered_to_date="int", expected_impressions_to_date="int",
        pacing_pct_to_date="pct", viewable_impressions="int", viewability_rate="pct", unique_reach="int",
        avg_frequency="dec", video_starts="int", video_completion_rate="pct", billable_impressions="int",
        cpm_usd="usd", revenue_usd="usd"), headers=dict(campaign_id="Campaign ID", cpm_usd="CPM (USD)",
                                                         revenue_usd="Revenue (USD)", adops_note="AdOps Note"))
    tr = last + 1
    ws.cell(row=tr, column=1, value="Total").font = BOLD
    for col in ["booked_impressions", "expected_impressions_month", "impressions", "viewable_impressions",
                "billable_impressions", "revenue_usd", "video_starts"]:
        L = get_column_letter(ccols.index(col) + 1)
        c = ws.cell(row=tr, column=ccols.index(col) + 1, value=f"=SUM({L}{first}:{L}{last})")
        c.font, c.number_format = BOLD, FMT["usd" if col == "revenue_usd" else "int"]
    imp_L, exp_L, v_L = (get_column_letter(ccols.index(x) + 1) for x in
                         ["impressions", "expected_impressions_month", "viewable_impressions"])
    c = ws.cell(row=tr, column=ccols.index("delivery_pct_month") + 1, value=f"=IFERROR({imp_L}{tr}/{exp_L}{tr},0)")
    c.font, c.number_format = BOLD, FMT["pct"]
    c = ws.cell(row=tr, column=ccols.index("viewability_rate") + 1, value=f"=IFERROR({v_L}{tr}/{imp_L}{tr},0)")
    c.font, c.number_format = BOLD, FMT["pct"]
    st = get_column_letter(ccols.index("pacing_status") + 1)
    rng = f"{st}{first}:{st}{last}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'OR({st}{first}="UNDER-PACING",{st}{first}="NO DELIVERY")'], fill=RED_FILL))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{st}{first}="OVER-PACING"'], fill=AMBER_FILL))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{st}{first}="ON PACE"'], fill=GREEN_FILL))
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(ccols))}{last}"
    autosize(ws)
    img = XLImage(str(charts["pacing"])); img.width, img.height = 900, 460
    ws.add_image(img, f"A{tr + 3}")

    # Placement Detail -------------------------------------------------------------
    ws = wb.create_sheet("Placement Detail")
    header_block(ws, "Placement Detail", "Inventory view. New inventory = placement launched this month (Product Mgmt catalog).")
    first, last = write_table(ws, placements, 4, formats=dict(
        launch_date="date", ad_requests="int", ad_server_filled="int", fill_rate="pct", telemetry_impressions="int",
        paid_impressions="int", house_impressions="int", sell_through_rate="pct", viewability_rate="pct",
        paid_unique_reach="int", video_completion_rate="pct", telemetry_vs_adserver_pct="pts"))
    vL = get_column_letter(list(placements.columns).index("viewability_rate") + 1)
    ws.conditional_formatting.add(f"{vL}{first}:{vL}{last}", CellIsRule(operator="lessThan", formula=["0.7"], fill=RED_FILL))
    nL = get_column_letter(list(placements.columns).index("is_new_inventory") + 1)
    ws.conditional_formatting.add(f"A{first}:{get_column_letter(len(placements.columns))}{last}",
                                  FormulaRule(formula=[f"${nL}{first}=TRUE"], fill=GREY_FILL))
    ws.freeze_panes = "D5"
    autosize(ws)
    img = XLImage(str(charts["viewability"])); img.width, img.height = 900, 400
    ws.add_image(img, f"A{last + 3}")

    # Frequency --------------------------------------------------------------------
    ws = wb.create_sheet("Frequency")
    header_block(ws, "Frequency Distribution", "Paid in-flight impressions per player in the month, by title.")
    write_table(ws, freq, 4, formats=dict(players="int", impressions="int", share_of_reach="pct",
                                          share_of_impressions="pct", cumulative_share_of_reach="pct"))
    autosize(ws)

    # Daily ------------------------------------------------------------------------
    ws = wb.create_sheet("Daily Trend")
    header_block(ws, "Daily Trend", "Paid in-flight impressions by title and day.")
    write_table(ws, daily, 4, formats=dict(event_date="date", paid_impressions="int", viewable_impressions="int",
                                           daily_unique_reach="int", paid_impressions_7d_avg="int"))
    ws.freeze_panes = "A5"
    autosize(ws)

    # Data Quality -----------------------------------------------------------------
    ws = wb.create_sheet("Data Quality")
    header_block(ws, "Data Quality & Automated Checks", "Error-severity failures block publication; warnings are listed for follow-up.")
    _, last = write_table(ws, dq, 4, formats=dict(value="int"))
    t = pd.DataFrame([{k: v for k, v in x.items() if k != "sample"} for x in tests])
    first, last2 = write_table(ws, t, last + 3)
    sL = get_column_letter(list(t.columns).index("status") + 1)
    ws.conditional_formatting.add(f"{sL}{first}:{sL}{last2}", FormulaRule(formula=[f'{sL}{first}="PASS"'], fill=GREEN_FILL))
    ws.conditional_formatting.add(f"{sL}{first}:{sL}{last2}", FormulaRule(formula=[f'{sL}{first}="WARN"'], fill=AMBER_FILL))
    ws.conditional_formatting.add(f"{sL}{first}:{sL}{last2}", FormulaRule(formula=[f'{sL}{first}="FAIL"'], fill=RED_FILL))
    autosize(ws, max_w=70)

    # AdOps Scope ------------------------------------------------------------------
    ws = wb.create_sheet("AdOps Scope")
    header_block(ws, "AdOps Reporting Scope", f"Campaign list agreed with AdOps for {month_label} (config/adops_scope_{month}.csv).")
    write_table(ws, scope, 4)
    autosize(ws)

    # Definitions ------------------------------------------------------------------
    ws = wb.create_sheet("Definitions")
    header_block(ws, "KPI Definitions", "Source: config/kpi_definitions.yml. Every KPI maps to a documented SQL model.")
    write_table(ws, kpis, 4)
    for row in ws.iter_rows(min_row=5):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCD", [28, 60, 34, 60]):
        ws.column_dimensions[col].width = w

    path = out_dir / f"AS_Inventory_Report_{month}.xlsx"
    wb.save(path)
    return path
